import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def _clean_html(text):
    if not text:
        return ""
    return text.replace("<p>", "").replace("</p>", "").replace("<strong>", "").replace("</strong>", "").replace("<em>", "").replace("</em>", "").replace("<br>", "\n").replace("&nbsp;", " ")

def build_docx(resume_data: dict) -> bytes:
    doc = Document()

    # Set basic margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    personal_info = resume_data.get("personalInfo", {})
    name = personal_info.get("name", "Name")

    # Name
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(24)

    # Contact Info
    contact_parts = []
    if personal_info.get("email"): contact_parts.append(personal_info["email"])
    if personal_info.get("phone"): contact_parts.append(personal_info["phone"])
    if personal_info.get("location"): contact_parts.append(personal_info["location"])
    if personal_info.get("linkedin"): contact_parts.append(personal_info["linkedin"])

    if contact_parts:
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_run = contact_p.add_run(" | ".join(contact_parts))
        contact_run.font.size = Pt(10)

    # Summary
    if resume_data.get("summary"):
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(_clean_html(resume_data["summary"]))

    # Experience
    experiences = resume_data.get("workExperience", [])
    if experiences:
        doc.add_heading("Experience", level=1)
        for exp in experiences:
            p = doc.add_paragraph()
            title_run = p.add_run(exp.get("title", ""))
            title_run.bold = True

            if exp.get("company"):
                p.add_run(f" at {exp['company']}")

            if exp.get("years"):
                p.add_run(f" | {exp['years']}")

            for bullet in exp.get("description", []):
                doc.add_paragraph(_clean_html(bullet), style='List Bullet')

    # Education
    educations = resume_data.get("education", [])
    if educations:
        doc.add_heading("Education", level=1)
        for ed in educations:
            p = doc.add_paragraph()
            deg_run = p.add_run(ed.get("degree", ""))
            deg_run.bold = True

            if ed.get("institution"):
                p.add_run(f", {ed['institution']}")
            if ed.get("years"):
                p.add_run(f" | {ed['years']}")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()
